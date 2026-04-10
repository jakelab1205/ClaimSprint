import email as _email_lib
import html as _html_lib
import re
from io import BytesIO

from .pdf_utils import PDFExtractionError, extract_text_from_pdf


class ExtractionError(ValueError):
    pass


_SUPPORTED_EXTENSIONS = frozenset({
    ".pdf", ".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp",
    ".docx", ".xlsx", ".txt", ".eml",
})


def extract_text_from_file(file_obj, filename: str) -> str:
    """
    Extract text from an uploaded file, dispatching by extension.
    Raises ExtractionError if the file cannot be read or yields no text.
    """
    import os
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        return _extract_pdf(file_obj)
    if ext in (".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"):
        return _extract_image(file_obj, filename)
    if ext == ".docx":
        return _extract_docx(file_obj)
    if ext == ".xlsx":
        return _extract_xlsx(file_obj)
    if ext == ".txt":
        return _extract_txt(file_obj)
    if ext == ".eml":
        return _extract_eml(file_obj)

    raise ExtractionError(
        f"Unsupported file type '{ext}'. "
        "Supported: PDF, JPG, PNG, TIFF, BMP, DOCX, XLSX, TXT, EML."
    )


# --- format handlers ---


def _extract_pdf(file_obj) -> str:
    try:
        return extract_text_from_pdf(file_obj)
    except PDFExtractionError as exc:
        raise ExtractionError(str(exc)) from exc


def _extract_image(file_obj, filename: str) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise ExtractionError("OCR dependencies (Pillow/pytesseract) are not installed.") from exc

    try:
        image = Image.open(file_obj)
        text = pytesseract.image_to_string(image)
    except Exception as exc:
        raise ExtractionError(f"Could not OCR '{filename}': {exc}") from exc

    text = text.strip()
    if not text:
        raise ExtractionError(
            f"No text could be extracted from '{filename}'. "
            "The image may be blank or too low-resolution for OCR."
        )
    return text


def _extract_docx(file_obj) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("python-docx is not installed.") from exc

    try:
        doc = docx.Document(file_obj)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX file: {exc}") from exc

    if not text:
        raise ExtractionError("No text content found in the DOCX file.")
    return text


def _extract_xlsx(file_obj) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ExtractionError("openpyxl is not installed.") from exc

    try:
        wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                row_text = "\t".join(cells).strip()
                if row_text:
                    lines.append(row_text)
    except Exception as exc:
        raise ExtractionError(f"Could not read XLSX file: {exc}") from exc

    content_lines = [l for l in lines if not l.startswith("[Sheet:")]
    if not content_lines:
        raise ExtractionError("No content found in the XLSX file.")
    return "\n".join(lines).strip()


def _extract_txt(file_obj) -> str:
    try:
        raw = file_obj.read()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
    except Exception as exc:
        raise ExtractionError(f"Could not read text file: {exc}") from exc

    text = text.strip()
    if not text:
        raise ExtractionError("The text file is empty.")
    return text


def _extract_eml(file_obj) -> str:
    try:
        raw = file_obj.read()
        msg = _email_lib.message_from_bytes(raw)
    except Exception as exc:
        raise ExtractionError(f"Could not parse EML file: {exc}") from exc

    parts = []

    # Headers
    header_lines = []
    for header in ("Subject", "From", "Date"):
        value = msg.get(header, "").strip()
        if value:
            header_lines.append(f"{header}: {value}")
    if header_lines:
        parts.append("\n".join(header_lines))

    # Body: prefer text/plain; fall back to stripped text/html
    plain_parts = []
    html_fallback = []
    attachment_parts = []

    for part in msg.walk():
        disposition = str(part.get("Content-Disposition", ""))
        content_type = part.get_content_type()

        if "attachment" in disposition:
            filename = part.get_filename() or ""
            payload = part.get_payload(decode=True)
            if payload and filename:
                import os
                ext = os.path.splitext(filename.lower())[1]
                if ext in _SUPPORTED_EXTENSIONS:
                    try:
                        extracted = extract_text_from_file(BytesIO(payload), filename)
                        attachment_parts.append(f"[Attachment: {filename}]\n{extracted}")
                    except ExtractionError:
                        pass
        elif content_type == "text/plain":
            payload = part.get_payload(decode=True)
            if payload:
                charset = part.get_content_charset() or "utf-8"
                plain_parts.append(payload.decode(charset, errors="replace"))
        elif content_type == "text/html":
            payload = part.get_payload(decode=True)
            if payload:
                charset = part.get_content_charset() or "utf-8"
                html_body = payload.decode(charset, errors="replace")
                clean = re.sub(r"<[^>]+>", " ", html_body)
                clean = _html_lib.unescape(clean)
                clean = re.sub(r"\s+", " ", clean).strip()
                html_fallback.append(clean)

    body_text = "\n".join(plain_parts).strip() if plain_parts else " ".join(html_fallback).strip()
    if body_text:
        parts.append(body_text)

    if attachment_parts:
        parts.append("\n\n".join(attachment_parts))

    full_text = "\n\n".join(p for p in parts if p).strip()
    if not full_text:
        raise ExtractionError("No readable content found in the EML file.")
    return full_text
